from django import forms
from django.contrib import admin, messages
from django.db import transaction
from django.http import HttpResponseRedirect
from django.shortcuts import render
from django.urls import path, reverse
from django.utils.html import format_html

from unfold.admin import ModelAdmin, StackedInline

from openpyxl import load_workbook

from .models import (
    CallQueueItem,
    CallSession,
    CompanyScript,
    Contact,
    CustomVoice,
    SalesInsight,
    SystemSettings,
)
# ============================================================
# CALL QUEUE ADMIN
# ============================================================

@admin.register(CallQueueItem)
class CallQueueItemAdmin(ModelAdmin):
    """
    Telicall lead queue admin.

    Excel import:
        /admin/calls/callqueueitem/import-excel/

    Expected workbook columns:
        Name
        Phone Number
        Details

    Name and Phone Number are required.
    Details is optional.
    """

    # A tiny change_list.html template can use this URL to render
    # an "Import Excel" button. The importer also works directly
    # at /admin/calls/callqueueitem/import-excel/.
    change_list_template = "admin/calls/callqueueitem/change_list.html"

    list_display = (
        "id",
        "name",
        "phone_number",
        "status",
        "formatted_duration",
        "top_question",
        "updated_at",
    )

    list_editable = (
        "status",
    )

    list_filter = (
        "status",
        "created_at",
        "updated_at",
    )

    search_fields = (
        "name",
        "phone_number",
        "details",
        "top_question",
    )

    ordering = (
        "-updated_at",
    )

    actions = [
        "mark_as_pending",
        "mark_as_called",
        "mark_as_followup",
    ]

    # --------------------------------------------------------
    # CUSTOM ADMIN URL
    # --------------------------------------------------------

    def get_urls(self):
        default_urls = super().get_urls()

        custom_urls = [
            path(
                "import-excel/",
                self.admin_site.admin_view(self.import_excel_view),
                name="calls_callqueueitem_import_excel",
            ),
        ]

        return custom_urls + default_urls

    # --------------------------------------------------------
    # EXCEL HELPERS
    # --------------------------------------------------------

    @staticmethod
    def _normalise_header(value):
        """
        Convert headings such as:
            Phone Number
            phone_number
            PHONE-NUMBER
        into:
            phonenumber
        """
        if value is None:
            return ""

        return "".join(
            ch.lower()
            for ch in str(value).strip()
            if ch.isalnum()
        )

    @staticmethod
    def _clean_cell(value):
        if value is None:
            return ""

        # Excel often gives numeric phone cells as float values.
        if isinstance(value, float) and value.is_integer():
            return str(int(value)).strip()

        return str(value).strip()

    @classmethod
    def _normalise_phone(cls, value):
        """
        Keep a leading + when present and remove common visual separators.

        Examples:
            98765 43210       -> 9876543210
            98765-43210       -> 9876543210
            +91 98765 43210   -> +919876543210
            9876543210.0      -> 9876543210
        """
        raw = cls._clean_cell(value)

        if not raw:
            return ""

        if raw.endswith(".0"):
            possible_number = raw[:-2]
            if possible_number.replace("+", "", 1).isdigit():
                raw = possible_number

        has_plus = raw.startswith("+")

        digits = "".join(
            ch for ch in raw
            if ch.isdigit()
        )

        if not digits:
            return ""

        return f"+{digits}" if has_plus else digits

    @staticmethod
    def _find_column(header_map, aliases):
        for alias in aliases:
            if alias in header_map:
                return header_map[alias]
        return None

    # --------------------------------------------------------
    # EXCEL IMPORT VIEW
    # --------------------------------------------------------

    def import_excel_view(self, request):
        context = {
            **self.admin_site.each_context(request),
            "title": "Import Leads from Excel",
            "opts": self.model._meta,
            "has_view_permission": self.has_view_permission(request),
            "has_add_permission": self.has_add_permission(request),
            "changelist_url": reverse(
                "admin:calls_callqueueitem_changelist"
            ),
        }

        if request.method != "POST":
            return render(
                request,
                "admin/calls/callqueueitem/import_excel.html",
                context,
            )

        excel_file = request.FILES.get("excel_file")

        if not excel_file:
            messages.error(
                request,
                "Please choose an Excel .xlsx file.",
            )
            return render(
                request,
                "admin/calls/callqueueitem/import_excel.html",
                context,
            )

        filename = str(
            getattr(excel_file, "name", "") or ""
        ).lower()

        if not filename.endswith(".xlsx"):
            messages.error(
                request,
                "Invalid file type. Please upload an .xlsx Excel file.",
            )
            return render(
                request,
                "admin/calls/callqueueitem/import_excel.html",
                context,
            )

        try:
            workbook = load_workbook(
                excel_file,
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            messages.error(
                request,
                f"Could not read the Excel file: {exc}",
            )
            return render(
                request,
                "admin/calls/callqueueitem/import_excel.html",
                context,
            )

        try:
            worksheet = workbook.active

            rows = worksheet.iter_rows(
                values_only=True
            )

            try:
                header_row = next(rows)
            except StopIteration:
                messages.error(
                    request,
                    "The Excel file is empty.",
                )
                return render(
                    request,
                    "admin/calls/callqueueitem/import_excel.html",
                    context,
                )

            header_map = {}

            for index, heading in enumerate(header_row):
                normalised = self._normalise_header(
                    heading
                )

                if normalised:
                    header_map[normalised] = index

            # Accepted Excel heading variants.
            name_index = self._find_column(
                header_map,
                {
                    "name",
                    "customername",
                    "leadname",
                    "fullname",
                },
            )

            phone_index = self._find_column(
                header_map,
                {
                    "phone",
                    "phonenumber",
                    "mobile",
                    "mobilenumber",
                    "contactnumber",
                    "contactphone",
                    "telephone",
                },
            )

            details_index = self._find_column(
                header_map,
                {
                    "details",
                    "detail",
                    "notes",
                    "note",
                    "description",
                    "remarks",
                    "remark",
                    "leaddetails",
                },
            )

            missing = []

            if name_index is None:
                missing.append("Name")

            if phone_index is None:
                missing.append("Phone Number")

            if missing:
                messages.error(
                    request,
                    (
                        "Required Excel column(s) missing: "
                        + ", ".join(missing)
                        + ". Required headings are Name and Phone Number."
                    ),
                )

                return render(
                    request,
                    "admin/calls/callqueueitem/import_excel.html",
                    context,
                )

            imported = 0
            duplicates = 0
            invalid = 0
            total_rows = 0

            # Prevent duplicates both against the database and within
            # the same workbook.
            existing_phones = set(
                str(phone).strip()
                for phone in CallQueueItem.objects.values_list(
                    "phone_number",
                    flat=True,
                )
                if phone
            )

            seen_in_file = set()

            new_items = []

            for excel_row_number, row in enumerate(
                rows,
                start=2,
            ):
                # Ignore completely blank rows.
                if not row or not any(
                    value not in (None, "")
                    for value in row
                ):
                    continue

                total_rows += 1

                try:
                    name_value = (
                        row[name_index]
                        if name_index < len(row)
                        else None
                    )

                    phone_value = (
                        row[phone_index]
                        if phone_index < len(row)
                        else None
                    )

                    details_value = (
                        row[details_index]
                        if (
                            details_index is not None
                            and details_index < len(row)
                        )
                        else None
                    )

                    name = self._clean_cell(
                        name_value
                    )

                    phone = self._normalise_phone(
                        phone_value
                    )

                    details = self._clean_cell(
                        details_value
                    )

                    # Required fields.
                    if not name or not phone:
                        invalid += 1
                        continue

                    # Basic sanity check. This remains intentionally broad
                    # because Telicall may later support multiple countries.
                    digit_count = len(
                        phone.lstrip("+")
                    )

                    if (
                        digit_count < 7
                        or digit_count > 15
                    ):
                        invalid += 1
                        continue

                    if (
                        phone in existing_phones
                        or phone in seen_in_file
                    ):
                        duplicates += 1
                        continue

                    new_items.append(
                        CallQueueItem(
                            name=name,
                            phone_number=phone,
                            details=details,
                            status="PENDING",
                        )
                    )

                    seen_in_file.add(phone)

                except Exception:
                    invalid += 1
                    continue

            if new_items:
                with transaction.atomic():
                    CallQueueItem.objects.bulk_create(
                        new_items,
                        batch_size=500,
                    )

                imported = len(new_items)

            messages.success(
                request,
                (
                    "Excel import completed. "
                    f"Rows checked: {total_rows} | "
                    f"Imported: {imported} | "
                    f"Duplicates skipped: {duplicates} | "
                    f"Invalid rows skipped: {invalid}"
                ),
            )

            return HttpResponseRedirect(
                reverse(
                    "admin:calls_callqueueitem_changelist"
                )
            )

        finally:
            try:
                workbook.close()
            except Exception:
                pass

    # --------------------------------------------------------
    # DISPLAY / BULK ACTIONS
    # --------------------------------------------------------

    @admin.display(description="Call Duration")
    def formatted_duration(self, obj):
        try:
            duration = int(
                obj.call_duration_seconds
            )
        except (ValueError, TypeError):
            duration = 0

        minutes, seconds = divmod(
            duration,
            60,
        )

        return f"{minutes}m {seconds}s"

    @admin.action(description="Mark selected leads as PENDING")
    def mark_as_pending(self, request, queryset):
        updated = queryset.update(
            status="PENDING"
        )

        self.message_user(
            request,
            f"{updated} lead(s) marked as PENDING.",
            level=messages.SUCCESS,
        )

    @admin.action(description="Mark selected leads as CALLED")
    def mark_as_called(self, request, queryset):
        updated = queryset.update(
            status="CALLED"
        )

        self.message_user(
            request,
            f"{updated} lead(s) marked as CALLED.",
            level=messages.SUCCESS,
        )

    @admin.action(description="Mark selected leads as FOLLOW_UP")
    def mark_as_followup(self, request, queryset):
        updated = queryset.update(
            status="FOLLOW_UP"
        )

        self.message_user(
            request,
            f"{updated} lead(s) marked as FOLLOW_UP.",
            level=messages.SUCCESS,
        )


# ============================================================
# COMPANY SCRIPT ADMIN
# ============================================================

@admin.register(CompanyScript)
class CompanyScriptAdmin(ModelAdmin):
    list_display = (
        "company_name",
        "bot_name",
        "tagline",
        "active_status_badge",
        "updated_at",
    )

    list_filter = (
        "is_active",
    )

    search_fields = (
        "company_name",
        "company_description",
        "products_services",
        "target_audience",
        "frequently_asked_questions",
    )

    readonly_fields = (
        "updated_at",
    )

    fieldsets = (
        (
            "Company Identity",
            {
                "fields": (
                    "company_name",
                    "bot_name",
                    "tagline",
                    "company_description",
                    "is_active",
                ),
            },
        ),
        (
            "Products & Audience",
            {
                "fields": (
                    "target_audience",
                    "products_services",
                    "key_benefits",
                    "course_durations",
                    "pricing_information",
                ),
            },
        ),
        (
            "Contact & FAQ",
            {
                "fields": (
                    "working_hours",
                    "contact_information",
                    "frequently_asked_questions",
                ),
            },
        ),
        (
            "AI Sales Configuration",
            {
                "fields": (
                    "sales_objective",
                    "data_to_collect",
                    "ai_rules",
                ),
            },
        ),
        (
            "Call Opening & Closing",
            {
                "fields": (
                    "opening_greeting",
                    "closing_statement",
                ),
            },
        ),
        (
            "Legacy / Additional Information",
            {
                "classes": (
                    "collapse",
                ),
                "fields": (
                    "company_details",
                    "updated_at",
                ),
            },
        ),
    )

    @admin.display(description="Active Status")
    def active_status_badge(self, obj):
        if obj.is_active:
            return format_html(
                (
                    '<span style="background-color: #d1fae5; '
                    'color: #065f46; '
                    'padding: 4px 10px; '
                    'border-radius: 12px; '
                    'font-weight: 600; '
                    'font-size: 11px;">'
                    "Active"
                    "</span>"
                )
            )

        return format_html(
            (
                '<span style="background-color: #f3f4f6; '
                'color: #374151; '
                'padding: 4px 10px; '
                'border-radius: 12px; '
                'font-weight: 600; '
                'font-size: 11px;">'
                "Inactive"
                "</span>"
            )
        )


# ============================================================
# SALES INSIGHT INLINE
# ============================================================

class SalesInsightInline(StackedInline):
    model = SalesInsight

    extra = 0

    readonly_fields = (
        "processed_at",
    )

    can_delete = False


# ============================================================
# CALL SESSION ADMIN
# ============================================================

@admin.register(CallSession)
class CallSessionAdmin(ModelAdmin):
    list_display = (
        "id",
        "contact",
        "formatted_duration",
        "status_badge",
        "audio_players",
        "created_at",
    )

    list_filter = (
        "status",
        "created_at",
    )

    search_fields = (
        "contact__phone_number",
        "contact__name",
    )

    readonly_fields = (
        "audio_players",
    )

    ordering = (
        "-created_at",
    )

    inlines = [
        SalesInsightInline,
    ]

    @admin.display(description="Duration")
    def formatted_duration(self, obj):
        try:
            duration = int(
                obj.duration_seconds
            )
        except (ValueError, TypeError):
            duration = 0

        minutes, seconds = divmod(
            duration,
            60,
        )

        return f"{minutes}m {seconds}s"

    @admin.display(description="Status")
    def status_badge(self, obj):
        colors = {
            "queued": "#d97706",
            "active": "#3b82f6",
            "completed": "#10b981",
            "failed": "#ef4444",
        }

        status = str(
            obj.status
        ).lower()

        bg_color = colors.get(
            status,
            "#6b7280",
        )

        return format_html(
            (
                '<span style="background-color: {}; '
                'color: #ffffff; '
                'padding: 3px 8px; '
                'border-radius: 6px; '
                'font-size: 11px; '
                'font-weight: 600; '
                'text-transform: uppercase;">'
                "{}"
                "</span>"
            ),
            bg_color,
            obj.status,
        )

    @admin.display(description="Audio Recordings")
    def audio_players(self, obj):
        html = []

        # --------------------------------------------
        # CUSTOMER AUDIO
        # --------------------------------------------

        if obj.recording_file:
            try:
                customer_url = (
                    obj.recording_file.url
                )

                html.append(
                    (
                        '<div style="margin-bottom: 8px;">'
                        '<span style="font-size: 11px; '
                        'color: #4b5563; '
                        'font-weight: 600;">'
                        "Customer Audio:"
                        "</span><br>"
                        '<audio controls preload="none" '
                        'style="height: 30px; '
                        'max-width: 220px;">'
                        f'<source src="{customer_url}" '
                        'type="audio/wav">'
                        "</audio>"
                        "</div>"
                    )
                )

            except Exception:
                html.append(
                    (
                        '<div style="margin-bottom: 8px;">'
                        '<span style="color: #dc2626; '
                        'font-size: 11px;">'
                        "Customer recording unavailable"
                        "</span>"
                        "</div>"
                    )
                )

        # --------------------------------------------
        # AI AUDIO
        # --------------------------------------------

        if obj.ai_recording_file:
            try:
                ai_url = (
                    obj.ai_recording_file.url
                )

                html.append(
                    (
                        "<div>"
                        '<span style="font-size: 11px; '
                        'color: #4b5563; '
                        'font-weight: 600;">'
                        "AI Response:"
                        "</span><br>"
                        '<audio controls preload="none" '
                        'style="height: 30px; '
                        'max-width: 220px;">'
                        f'<source src="{ai_url}" '
                        'type="audio/wav">'
                        "</audio>"
                        "</div>"
                    )
                )

            except Exception:
                html.append(
                    (
                        "<div>"
                        '<span style="color: #dc2626; '
                        'font-size: 11px;">'
                        "AI recording unavailable"
                        "</span>"
                        "</div>"
                    )
                )

        if html:
            return format_html(
                "".join(html)
            )

        return format_html(
            (
                '<span style="color: #9ca3af; '
                'font-style: italic;">'
                "No Recordings"
                "</span>"
            )
        )


# ============================================================
# CONTACT ADMIN
# ============================================================

@admin.register(Contact)
class ContactAdmin(ModelAdmin):
    list_display = (
        "name",
        "phone_number",
        "uploaded_at",
    )

    search_fields = (
        "name",
        "phone_number",
    )

    ordering = (
        "-uploaded_at",
    )


# ============================================================
# SALES INSIGHT ADMIN
# ============================================================

@admin.register(SalesInsight)
class SalesInsightAdmin(ModelAdmin):
    list_display = (
        "call_session",
        "followup_badge",
        "processed_at",
    )

    list_filter = (
        "needs_followup",
        "processed_at",
    )

    readonly_fields = (
        "processed_at",
    )

    @admin.display(description="Needs Followup")
    def followup_badge(self, obj):
        if obj.needs_followup:
            return format_html(
                (
                    '<span style="background-color: #fef2f2; '
                    'color: #991b1b; '
                    'border: 1px solid #fecaca; '
                    'padding: 3px 8px; '
                    'border-radius: 12px; '
                    'font-weight: 600; '
                    'font-size: 11px;">'
                    "Action Required"
                    "</span>"
                )
            )

        return format_html(
            (
                '<span style="background-color: #f0fdf4; '
                'color: #166534; '
                'border: 1px solid #bbf7d0; '
                'padding: 3px 8px; '
                'border-radius: 12px; '
                'font-weight: 600; '
                'font-size: 11px;">'
                "Resolved"
                "</span>"
            )
        )

# ============================================================
# CUSTOM VOICE ADMIN
# ============================================================

@admin.register(CustomVoice)
class CustomVoiceAdmin(ModelAdmin):

    list_display = (
        "name",
        "language",
        "status",
        "sample_count",
        "created_at",
        "updated_at",
    )

    list_filter = (
        "status",
        "language",
        "created_at",
    )

    search_fields = (
        "name",
        "description",
    )

    readonly_fields = (
        "created_at",
        "updated_at",
        "sample_1_preview",
        "sample_2_preview",
        "sample_3_preview",
    )

    fieldsets = (
        (
            "Voice Information",
            {
                "fields": (
                    "name",
                    "language",
                    "status",
                    "description",
                ),
            },
        ),

        (
            "Voice Recording Samples",
            {
                "fields": (
                    "sample_1",
                    "sample_1_preview",
                    "sample_2",
                    "sample_2_preview",
                    "sample_3",
                    "sample_3_preview",
                ),
                "description": (
                    "Choose an existing WAV file or use the "
                    "Record Voice button beside each sample."
                ),
            },
        ),

        (
            "System Information",
            {
                "classes": ("collapse",),
                "fields": (
                    "created_at",
                    "updated_at",
                ),
            },
        ),
    )

    class Media:
        js = (
            "calls/admin/custom_voice_recorder.js",
        )

    @admin.display(description="Samples")
    def sample_count(self, obj):
        count = sum(
            bool(sample)
            for sample in (
                obj.sample_1,
                obj.sample_2,
                obj.sample_3,
            )
        )

        return f"{count}/3"

    def _audio_preview(self, audio_file):
        if not audio_file:
            return format_html(
                '<span style="color:#9ca3af;">No recording</span>'
            )

        try:
            return format_html(
                '<audio controls preload="metadata" style="width:320px;">'
                '<source src="{}" type="audio/wav">'
                "</audio>",
                audio_file.url,
            )
        except Exception:
            return "Audio unavailable"

    @admin.display(description="Preview")
    def sample_1_preview(self, obj):
        return self._audio_preview(obj.sample_1)

    @admin.display(description="Preview")
    def sample_2_preview(self, obj):
        return self._audio_preview(obj.sample_2)

    @admin.display(description="Preview")
    def sample_3_preview(self, obj):
        return self._audio_preview(obj.sample_3)
# ============================================================
# TELICALL SYSTEM SETTINGS ADMIN
# ============================================================


class SystemSettingsAdminForm(forms.ModelForm):
    """
    Validates the selected AI voice configuration.

    - Windows SAPI requires a SAPI voice unless system default is enabled.
    - Custom Voice requires a saved CustomVoice record with at least one
      reference sample.
    """

    class Meta:
        model = SystemSettings
        fields = "__all__"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)

        # Show only usable custom voices in the dropdown.
        if "custom_voice" in self.fields:
            self.fields["custom_voice"].queryset = (
                CustomVoice.objects
                .filter(status="ready")
                .order_by("name")
            )

            self.fields["custom_voice"].help_text = (
                "Select a READY custom voice. Create/edit voices from "
                "Custom Voices in Django Admin."
            )

    def clean(self):
        cleaned = super().clean()

        engine = str(
            cleaned.get("tts_engine") or "sapi"
        ).strip().lower()

        use_default = bool(
            cleaned.get("use_system_default_voice")
        )

        sapi_voice = str(
            cleaned.get("tts_voice_name") or ""
        ).strip()

        custom_voice = cleaned.get("custom_voice")

        if engine == "sapi":
            if not use_default and not sapi_voice:
                self.add_error(
                    "tts_voice_name",
                    (
                        "Select a Windows SAPI voice or enable "
                        "'Use system default voice'."
                    ),
                )

        elif engine == "custom":
            if custom_voice is None:
                self.add_error(
                    "custom_voice",
                    (
                        "Select a custom voice before saving "
                        "Custom Voice as the TTS engine."
                    ),
                )
            else:
                has_sample = any(
                    bool(getattr(custom_voice, field_name, None))
                    for field_name in (
                        "sample_1",
                        "sample_2",
                        "sample_3",
                    )
                )

                if not has_sample:
                    self.add_error(
                        "custom_voice",
                        (
                            "The selected custom voice has no reference "
                            "recording. Upload or record at least one sample."
                        ),
                    )

                if getattr(custom_voice, "status", "") != "ready":
                    self.add_error(
                        "custom_voice",
                        (
                            "The selected custom voice must have status READY "
                            "before it can be used for TTS."
                        ),
                    )

        return cleaned


@admin.register(SystemSettings)
class SystemSettingsAdmin(ModelAdmin):

    form = SystemSettingsAdminForm

    list_display = (
        "voice_status",
        "tts_rate",
        "tts_volume",
        "retention_period",
        "cleanup_status",
        "updated_at",
    )

    readonly_fields = (
        "updated_at",
    )

    fieldsets = (
        (
            "AI Voice Settings",
            {
                "fields": (
                    "tts_engine",
                    "use_system_default_voice",
                    "tts_voice_name",
                    "custom_voice",
                    "tts_rate",
                    "tts_volume",
                ),
                "description": (
                    "To use a cloned/custom voice: set Voice Engine to "
                    "'Custom Voice', select a READY voice from Custom Voice, "
                    "then save. To use Windows speech: select 'Windows SAPI'."
                ),
            },
        ),

        (
            "Call Recording Cleanup",
            {
                "fields": (
                    "automatic_recording_cleanup",
                    "recording_retention_days",
                ),
            },
        ),

        (
            "System Information",
            {
                "classes": ("collapse",),
                "fields": (
                    "updated_at",
                ),
            },
        ),
    )

    @admin.display(description="AI Voice")
    def voice_status(self, obj):
        if obj.tts_engine == "custom":
            if obj.custom_voice:
                return format_html(
                    '<span style="background:#ede9fe;'
                    'color:#5b21b6;padding:4px 10px;'
                    'border-radius:12px;font-weight:600;">'
                    'Custom: {}'
                    '</span>',
                    obj.custom_voice.name,
                )

            return format_html(
                '<span style="background:#fef2f2;'
                'color:#991b1b;padding:4px 10px;'
                'border-radius:12px;font-weight:600;">'
                'Custom Voice - Not Selected'
                '</span>'
            )

        if obj.use_system_default_voice:
            return format_html(
                '<span style="background:#e0f2fe;'
                'color:#075985;padding:4px 10px;'
                'border-radius:12px;font-weight:600;">'
                'Windows Default'
                '</span>'
            )

        return format_html(
            '<span style="background:#e0f2fe;'
            'color:#075985;padding:4px 10px;'
            'border-radius:12px;font-weight:600;">'
            '{}'
            '</span>',
            obj.tts_voice_name,
        )

    @admin.display(description="Retention Period")
    def retention_period(self, obj):
        return (
            f"{obj.recording_retention_days} "
            f"{'Day' if obj.recording_retention_days == 1 else 'Days'}"
        )

    @admin.display(description="Automatic Cleanup")
    def cleanup_status(self, obj):
        if obj.automatic_recording_cleanup:
            return format_html(
                '<span style="background:#d1fae5;'
                'color:#065f46;padding:4px 10px;'
                'border-radius:12px;font-weight:600;">'
                "Enabled"
                "</span>"
            )

        return format_html(
            '<span style="background:#fef2f2;'
            'color:#991b1b;padding:4px 10px;'
            'border-radius:12px;font-weight:600;">'
            "Disabled"
            "</span>"
        )

    def has_add_permission(self, request):
        return not SystemSettings.objects.exists()

    def has_delete_permission(self, request, obj=None):
        return False
